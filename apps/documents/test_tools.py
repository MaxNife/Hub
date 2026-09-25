"""Engine tests on real files. Run: python -m unittest apps/documents/test_tools.py"""

import importlib.util
import io
import pathlib
import sys
import unittest
import zipfile

import pymupdf

here = pathlib.Path(__file__).parent
sys.path.insert(0, str(here))
spec = importlib.util.spec_from_file_location("doc_tools", here / "tools.py")
tools = importlib.util.module_from_spec(spec)
sys.modules["doc_tools"] = tools  # dataclasses look their module up
spec.loader.exec_module(tools)


def make_pdf(name: str, pages: int, label: str = "Page") -> tools.File:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 100), f"{label} {i + 1}", fontsize=24)
        page.insert_text((72, 140), "Confidential salary 90000", fontsize=12)
    return tools.File(name, doc.tobytes())


def png(name: str, w: int, h: int, rgb=None) -> tools.File:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), 0)
    if rgb:
        pix.set_rect(pix.irect, rgb)
    return tools.File(name, pix.tobytes("png"))


def opened(f: tools.File) -> pymupdf.Document:
    return pymupdf.open(stream=f.data, filetype="pdf")


def text_of(f, i=0):
    with opened(f) as d:
        return d[i].get_text()


def names_in(z: tools.File) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(z.data)) as zf:
        return zf.namelist()


class ToolsTest(unittest.TestCase):
    def setUp(self):
        self.a = make_pdf("a.pdf", 3, "Alpha")
        self.b = make_pdf("b.pdf", 2, "Beta")

    def run_tool(self, tool, files, **opts):
        return tools.run(tool, files, opts)

    def pages(self, f):
        with opened(f) as d:
            return d.page_count

    def test_parse_pages(self):
        self.assertEqual(tools.parse_pages("1-3, 5", 6), [0, 1, 2, 4])
        self.assertEqual(tools.parse_pages("4-", 6), [3, 4, 5])
        self.assertEqual(tools.parse_pages("all", 3), [0, 1, 2])
        with self.assertRaises(tools.ToolError):
            tools.parse_pages("7", 6)
        with self.assertRaises(tools.ToolError):
            tools.parse_pages("x", 6)

    def test_merge_includes_images(self):
        [out] = self.run_tool("merge", [self.a, self.b, png("pic.png", 40, 30, (255, 150, 0))])
        self.assertEqual(self.pages(out), 6)
        self.assertIn("Beta 1", text_of(out, 3))

    def test_split_modes(self):
        [z] = self.run_tool("split", [self.a], mode="ranges", ranges="1-2, 3")
        self.assertEqual(sorted(names_in(z)), ["a pages 1-2.pdf", "a pages 3.pdf"])
        [one] = tools.run("split", [self.a], {"mode": "extract", "pages": "2"})
        self.assertEqual(self.pages(one), 1)
        self.assertIn("Alpha 2", text_of(one))
        [z3] = tools.run("split", [self.a], {"mode": "every", "every": 1})
        self.assertEqual(len(names_in(z3)), 3)

    def test_remove_organize_rotate(self):
        [out] = self.run_tool("remove", [self.a], pages="2")
        self.assertEqual(self.pages(out), 2)
        self.assertIn("Alpha 3", text_of(out, 1))
        [org] = tools.run("organize", [self.a], {"pages": [{"index": 2, "rotate": 90}, {"index": 0}]})
        with opened(org) as d:
            self.assertEqual((d.page_count, d[0].rotation, d[1].rotation), (2, 90, 0))
            self.assertIn("Alpha 3", d[0].get_text())
        [rot] = tools.run("rotate", [self.a], {"angle": 180, "pages": "1"})
        with opened(rot) as d:
            self.assertEqual([p.rotation for p in d], [180, 0, 0])
        with self.assertRaises(tools.ToolError):
            tools.run("remove", [self.a], {"pages": "1-3"})

    def test_compress_never_grows(self):
        doc = pymupdf.open()
        page = doc.new_page()
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 1600, 1600), 0)
        for y in range(0, 1600, 40):  # noisy stripes so there is something to squeeze
            pix.set_rect(pymupdf.IRect(0, y, 1600, y + 20), ((y * 7) % 255, (y * 3) % 255, 90))
        page.insert_image(pymupdf.Rect(0, 0, 200, 200), pixmap=pix)
        big = tools.File("big.pdf", doc.tobytes())
        [out] = self.run_tool("compress", [big], level="extreme")
        self.assertLessEqual(out.size, big.size)
        self.assertEqual(out.name, "big compressed.pdf")

    def test_images_to_pdf(self):
        scan = png("scan.png", 300, 400)
        [out] = self.run_tool("images-to-pdf", [scan, scan], size="a4", margin="small")
        with opened(out) as d:
            self.assertEqual(d.page_count, 2)
            self.assertAlmostEqual(d[0].rect.width, 595, delta=1)

    def test_pdf_to_images_text_word(self):
        [z] = self.run_tool("pdf-to-images", [self.a], format="jpg", dpi=72)
        self.assertEqual(len(names_in(z)), 3)
        [txt] = tools.run("pdf-to-text", [self.a], {})
        self.assertIn("--- Page 2 ---\nAlpha 2", txt.data.decode())
        if tools.DocxConverter:
            [docx] = tools.run("pdf-to-word", [self.a], {})
            self.assertEqual(docx.suffix, ".docx")
            with zipfile.ZipFile(io.BytesIO(docx.data)) as zf:
                self.assertIn("Alpha 1", zf.read("word/document.xml").decode())

    @unittest.skipUnless(tools.soffice(), "LibreOffice not installed")
    def test_office_to_pdf(self):
        from docx import Document  # comes with pdf2docx
        d = Document()
        d.add_heading("Quarterly report", 0)
        d.add_paragraph("Revenue went up.")
        buf = io.BytesIO()
        d.save(buf)
        [out] = self.run_tool("office-to-pdf", [tools.File("report.docx", buf.getvalue())])
        self.assertEqual(out.name, "report.pdf")
        self.assertIn("Quarterly report", text_of(out))

    def test_page_numbers_and_watermark(self):
        [out] = self.run_tool("page-numbers", [self.a], style="of", position="bottom-right")
        self.assertIn("Page 3 of 3", text_of(out, 2))
        [wm] = tools.run("watermark", [self.a], {"text": "DRAFT", "opacity": 0.2})
        self.assertIn("DRAFT", text_of(wm, 1))
        with self.assertRaises(tools.ToolError):
            tools.run("watermark", [self.a], {"text": " "})

    def test_edit_ops_and_redaction(self):
        import base64
        data = "data:image/png;base64," + base64.b64encode(png("sig.png", 60, 20).data).decode()
        ops = [
            {"page": 0, "type": "text", "x": 72, "y": 200, "w": 150, "h": 20, "text": "Approved by Hope", "size": 14, "color": "#1e7a48"},
            {"page": 0, "type": "rect", "x": 60, "y": 120, "w": 300, "h": 30, "fill": "#ffffff"},
            {"page": 0, "type": "ink", "points": [[72, 300], [120, 320], [160, 300]], "color": "#000", "width": 2},
            {"page": 1, "type": "image", "x": 72, "y": 400, "w": 120, "h": 40, "data": data},
            {"page": 1, "type": "redact", "x": 60, "y": 125, "w": 300, "h": 20},
        ]
        [out] = self.run_tool("edit", [self.a], ops=ops)
        with opened(out) as d:
            self.assertIn("Approved by Hope", d[0].get_text())
            self.assertNotIn("salary", d[1].get_text())  # redaction really removes text
            self.assertIn("salary", d[2].get_text())
            self.assertEqual(len(d[1].get_images()), 1)

    def test_edit_rotated_page_keeps_text_upright(self):
        doc = opened(self.a)
        doc[0].set_rotation(90)
        rot = tools.File("rot.pdf", doc.tobytes())
        [out] = self.run_tool("edit", [rot], ops=[{"page": 0, "type": "text", "x": 50, "y": 50, "w": 200, "h": 30, "text": "Hello", "size": 16}])
        with opened(out) as d:
            words = d[0].get_text("words")  # unrotated coordinates
            hello = [w for w in words if w[4] == "Hello"][0]
            seen = pymupdf.Rect(hello[:4]) * d[0].rotation_matrix
            self.assertLess(seen.x0, 70)  # where it was placed on the visible page
            self.assertLess(seen.y0, 70)
            self.assertGreater(seen.width, seen.height)  # reads left to right, not sideways

    def test_protect_then_unlock(self):
        [locked] = self.run_tool("protect", [self.a], password="s3cret")
        with opened(locked) as d:
            self.assertTrue(d.needs_pass)
        with self.assertRaises(tools.ToolError):
            tools.run("merge", [locked, self.b], {})
        with self.assertRaises(tools.ToolError):
            tools.run("unlock", [locked], {"password": "nope"})
        [free] = tools.run("unlock", [locked], {"password": "s3cret"})
        with opened(free) as d:
            self.assertFalse(d.needs_pass)
            self.assertIn("Alpha 1", d[0].get_text())

    def test_nothing_touches_the_disk(self):
        """Apart from LibreOffice's scratch folder, tools never write files."""
        import builtins
        import unittest.mock as mock
        real_open = builtins.open
        def guarded(file, mode="r", *a, **k):
            if any(c in mode for c in "wax+"):
                raise AssertionError(f"wrote {file}")
            return real_open(file, mode, *a, **k)
        real_save = pymupdf.Document.save
        def save(doc, target, *a, **k):  # tobytes() saves into a BytesIO, which is fine
            if isinstance(target, (str, pathlib.PurePath)):
                raise AssertionError(f"saved {target}")
            return real_save(doc, target, *a, **k)
        with mock.patch("builtins.open", guarded), mock.patch.object(pymupdf.Document, "save", save):
            for tool, files, opts in [("merge", [self.a, self.b], {}), ("split", [self.a], {"mode": "every", "every": 1}),
                                      ("compress", [self.a], {}), ("pdf-to-images", [self.a], {"dpi": 72}),
                                      ("pdf-to-text", [self.a], {}), ("protect", [self.a], {"password": "abcd"}),
                                      ("watermark", [self.a], {"text": "X"}), ("page-numbers", [self.a], {})]:
                self.assertTrue(tools.run(tool, files, opts)[0].data, tool)


if __name__ == "__main__":
    unittest.main()
